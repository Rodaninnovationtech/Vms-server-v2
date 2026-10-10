import os
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from django.conf import settings
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

# ---------------------------------------------------------------------
# shared look
# ---------------------------------------------------------------------

REPORT_TITLE = "Visitor / Contractor Check-in Report"
REPORT_TZ = ZoneInfo("Asia/Kolkata")      # times in the file are shown in this zone

PURPLE_HEX = "6A1B9A"                     # heading colour (Excel + PDF)
ZEBRA_HEX = "F5EEFB"                      # light purple for alternate rows
BORDER_HEX = "D9C7EA"


def logo_path():
    """Absolute path of the admin logo, or None if the file is missing."""
    path = os.path.join(settings.MEDIA_ROOT, settings.SUPER_ADMIN_LOGO)
    return path if os.path.isfile(path) else None


def fmt_dt(value):
    """ISO datetime string -> dd-mm-yyyy hh:mm AM/PM (in REPORT_TZ)."""
    if not value:
        return ""
    try:
        dt = datetime.fromisoformat(value)
    except ValueError:
        return str(value)
    if dt.tzinfo is not None:
        dt = dt.astimezone(REPORT_TZ)
    return dt.strftime("%d-%m-%Y %I:%M %p")


def build_subtitle(filters, total):
    parts = []
    f, t = filters.get("from_date"), filters.get("to_date")
    if f or t:
        parts.append(
            "Check-in: "
            f"{f.strftime('%d-%m-%Y') if f else '...'} to "
            f"{t.strftime('%d-%m-%Y') if t else '...'}"
        )
    if filters.get("search"):
        parts.append(f"Search: {filters['search']}")
    parts.append(f"Total records: {total}")
    return "   |   ".join(parts)


def _text(row, key):
    value = row.get(key)
    if key in ("check_in", "check_out"):
        return fmt_dt(value)
    return "" if value is None else str(value)


# ---------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------

# (header, row key, column width)
EXCEL_COLUMNS = [
    ("Visitor Name", "person_name", 24),
    ("Visitor Type", "visitor_type_name", 14),
    ("Identity Type", "identity_type", 16),
    ("Identity Number", "identity_number", 18),
    ("Contact", "phone_number", 16),
    ("Email", "email", 28),
    ("Company", "company", 22),
    ("Company Phone", "company_phone", 16),
    ("Site", "site_name", 20),
    ("Location", "location", 20),
    ("Pass No", "pass_no", 11),
    ("Key No", "key_no", 11),
    ("Vehicle", "vehicle_number", 14),
    ("Check In", "check_in", 20),
    ("Check Out", "check_out", 20),
    ("Status", "status", 14),
]


def build_excel(rows, subtitle):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.sheet_view.showGridLines = False

    ncols = len(EXCEL_COLUMNS)
    last_col = get_column_letter(ncols)

    purple = PatternFill("solid", fgColor=PURPLE_HEX)
    white = PatternFill("solid", fgColor="FFFFFF")
    zebra = PatternFill("solid", fgColor=ZEBRA_HEX)
    side = Side(style="thin", color=BORDER_HEX)
    border = Border(left=side, right=side, top=side, bottom=side)

    for i, (_, _, width) in enumerate(EXCEL_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # ---- heading band (rows 1-2): logo on the left, title on purple ----
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 8

    logo = logo_path()
    for r in (1, 2):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = white if (logo and c == 1) else purple

    if logo:
        with PILImage.open(logo) as im:
            w, h = im.size
        scale = min(70 / h, 165 / w)          # fit inside column A, rows 1-2
        img = XLImage(logo)
        img.width, img.height = int(w * scale), int(h * scale)
        ws.add_image(img, "A1")
        title_col = "B"
    else:
        title_col = "A"

    ws[f"{title_col}1"] = REPORT_TITLE
    ws[f"{title_col}1"].font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    ws[f"{title_col}1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(f"{title_col}1:{last_col}1")

    ws[f"{title_col}2"] = subtitle
    ws[f"{title_col}2"].font = Font(name="Calibri", size=10, color="FFFFFF")
    ws[f"{title_col}2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(f"{title_col}2:{last_col}2")

    # ---- column headings (row 4) ----
    header_row = 4
    ws.row_dimensions[header_row].height = 24
    for i, (header, _, _) in enumerate(EXCEL_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=i, value=header)
        cell.fill = purple
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    # ---- data ----
    for n, row in enumerate(rows):
        r = header_row + 1 + n
        for i, (_, key, _) in enumerate(EXCEL_COLUMNS, start=1):
            cell = ws.cell(row=r, column=i, value=_text(row, key))
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(vertical="center")
            cell.border = border
            if n % 2 == 1:
                cell.fill = zebra

    ws.freeze_panes = f"A{header_row + 1}"
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + len(rows)}"

    # printing
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"1:{header_row}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------

PURPLE = colors.HexColor(f"#{PURPLE_HEX}")
ZEBRA = colors.HexColor(f"#{ZEBRA_HEX}")
BORDER = colors.HexColor(f"#{BORDER_HEX}")

PAGE = landscape(A4)
MARGIN = 8 * mm
HEADER_H = 26 * mm

CELL = ParagraphStyle("cell", fontName="Helvetica", fontSize=6.5, leading=8)
HEAD = ParagraphStyle(
    "head", fontName="Helvetica-Bold", fontSize=7, leading=8.5, textColor=colors.white
)


def _p(text, style=CELL):
    return Paragraph(escape(text), style)


def _identity(r):
    return Paragraph(
        f"{escape(r.get('identity_type') or '')}<br/>{escape(r.get('identity_number') or '')}",
        CELL,
    )


def _pass_key(r):
    parts = []
    if r.get("pass_no"):
        parts.append(f"Pass: {escape(r['pass_no'])}")
    if r.get("key_no"):
        parts.append(f"Key: {escape(r['key_no'])}")
    return Paragraph("<br/>".join(parts), CELL)


# (header, cell builder, width in mm) - widths add up to the 281 mm printable width
PDF_COLUMNS = [
    ("Visitor Name", lambda r: _p(_text(r, "person_name")), 25),
    ("Type", lambda r: _p(_text(r, "visitor_type_name")), 16),
    ("Identity", _identity, 28),
    ("Contact", lambda r: _p(_text(r, "phone_number")), 21),
    ("Email", lambda r: _p(_text(r, "email")), 30),
    ("Company", lambda r: _p(_text(r, "company")), 24),
    ("Site", lambda r: _p(_text(r, "site_name")), 22),
    ("Location", lambda r: _p(_text(r, "location")), 20),
    ("Pass / Key", _pass_key, 16),
    ("Vehicle", lambda r: _p(_text(r, "vehicle_number")), 18),
    ("Check In", lambda r: _p(_text(r, "check_in")), 22),
    ("Check Out", lambda r: _p(_text(r, "check_out")), 22),
    ("Status", lambda r: _p(_text(r, "status")), 17),
]


def build_pdf(rows, subtitle):
    buf = BytesIO()
    generated = datetime.now(REPORT_TZ).strftime("%d-%m-%Y %I:%M %p")
    logo = logo_path()

    def decorate(canvas, doc):
        """Purple heading with logo - drawn on every page."""
        width, height = PAGE
        canvas.saveState()

        canvas.setFillColor(PURPLE)
        canvas.rect(0, height - HEADER_H, width, HEADER_H, stroke=0, fill=1)

        text_x = MARGIN
        if logo:
            img = ImageReader(logo)
            iw, ih = img.getSize()
            h = HEADER_H - 10 * mm
            w = iw * h / ih
            if w > 50 * mm:
                w = 50 * mm
                h = ih * w / iw
            pad = 1.5 * mm
            y0 = height - HEADER_H + (HEADER_H - h) / 2

            canvas.setFillColor(colors.white)
            canvas.roundRect(
                MARGIN, y0 - pad, w + 2 * pad, h + 2 * pad, 1.5 * mm, stroke=0, fill=1
            )
            canvas.drawImage(img, MARGIN + pad, y0, w, h, mask="auto")
            text_x = MARGIN + w + 2 * pad + 5 * mm

        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 16)
        canvas.drawString(text_x, height - 12 * mm, REPORT_TITLE)
        canvas.setFont("Helvetica", 8.5)
        canvas.drawString(text_x, height - 19 * mm, subtitle)
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(width - MARGIN, height - 12 * mm, f"Generated: {generated}")

        # footer
        canvas.setFillColor(colors.grey)
        canvas.setFont("Helvetica", 7)
        canvas.drawString(MARGIN, 6 * mm, REPORT_TITLE)
        canvas.drawRightString(width - MARGIN, 6 * mm, f"Page {doc.page}")

        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf,
        pagesize=PAGE,
        leftMargin=MARGIN,
        rightMargin=MARGIN,
        topMargin=HEADER_H + 5 * mm,
        bottomMargin=12 * mm,
        title=REPORT_TITLE,
    )

    if not rows:
        story = [Paragraph("No records found for the selected filters.", CELL)]
    else:
        data = [[Paragraph(escape(h), HEAD) for h, _, _ in PDF_COLUMNS]]
        for r in rows:
            data.append([build(r) for _, build, _ in PDF_COLUMNS])

        table = Table(
            data,
            colWidths=[w * mm for _, _, w in PDF_COLUMNS],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), PURPLE),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ZEBRA]),
                    ("GRID", (0, 0), (-1, -1), 0.25, BORDER),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story = [table]

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()